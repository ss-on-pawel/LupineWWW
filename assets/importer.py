import io
from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.exceptions import ValidationError

from locations.models import Location
from .models import Asset, AssetHistoryEntry
from .services import generate_unique_asset_barcode, record_asset_history


def parse_import_xlsx(file_obj):
    """
    Parse XLSX import file starting from row 2 (row 1 = header).
    Skips completely empty rows and rows where the first cell starts with '#'.
    Returns list of row dicts.
    """
    try:
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Nie można otworzyć pliku XLSX: {exc}") from exc

    ws = wb.active
    rows = []

    for i, row_cells in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        raw = list(row_cells)
        cells = [("" if c is None else str(c).strip()) for c in (raw + [""] * 5)[:5]]
        if not any(cells):
            continue
        if cells[0].startswith("#"):
            continue
        rows.append({
            "row_number": i + 2,
            "inventory_number": cells[0],
            "name": cells[1],
            "quantity": cells[2],
            "value": cells[3],
            "location_path": cells[4],
        })

    wb.close()
    return rows


def _validate_row(row):
    """Validate a single parsed row. Returns list of error strings (empty = valid)."""
    errors = []

    name = row.get("name", "")
    if not name:
        errors.append("Brak nazwy")
    elif len(name) > 255:
        errors.append("Nazwa za długa (maks. 255 znaków)")

    location_path = row.get("location_path", "")
    if not location_path:
        errors.append("Brak lokalizacji")
    else:
        segments = [s.strip() for s in location_path.split("/")]
        if any(s == "" for s in segments):
            errors.append("Niepoprawna lokalizacja (puste segmenty lub podwójny /)")

    quantity_str = row.get("quantity", "")
    if quantity_str:
        try:
            qty_float = float(quantity_str)
            if qty_float != int(qty_float):
                errors.append("Niepoprawna ilość (musi być liczbą całkowitą)")
            elif int(qty_float) <= 0:
                errors.append("Ilość musi być większa od zera")
        except (ValueError, TypeError):
            errors.append("Niepoprawna ilość (musi być liczbą całkowitą)")

    value_str = row.get("value", "")
    if value_str:
        try:
            Decimal(value_str.replace(",", "."))
        except InvalidOperation:
            errors.append("Niepoprawna wartość (musi być liczbą, np. 1500.00)")

    return errors


def resolve_location_path(root_location, path_str):
    """
    Resolve or create location hierarchy under root_location.
    Returns (location, new_location_ids) where new_location_ids is a set of
    PKs of locations created during this call.
    Raises ValidationError on empty path.
    """
    segments = [s.strip() for s in path_str.split("/") if s.strip()]
    if not segments:
        raise ValidationError("Ścieżka lokalizacji jest pusta.")

    current = root_location
    new_ids = set()
    for segment in segments:
        loc, created = Location.objects.get_or_create(
            parent=current,
            name=segment,
            defaults={"is_active": True},
        )
        if created:
            new_ids.add(loc.pk)
        current = loc

    return current, new_ids


def import_assets_from_rows(rows, *, root_location, asset_type_ref, operator=None):
    """
    Import assets with partial success.
    Valid rows are saved; invalid rows are skipped and reported.
    Returns dict: imported_count, created_locations_count, error_rows.
    """
    imported_count = 0
    created_location_ids = set()
    error_rows = []

    for row in rows:
        row_number = row["row_number"]

        row_errors = _validate_row(row)
        if row_errors:
            for err in row_errors:
                error_rows.append({"row": row_number, "message": err})
            continue

        quantity_str = row.get("quantity", "")
        quantity = int(float(quantity_str)) if quantity_str else 1

        value_str = row.get("value", "")
        purchase_value = Decimal(value_str.replace(",", ".")) if value_str else None

        try:
            barcode = generate_unique_asset_barcode(asset_type_ref=asset_type_ref)
        except ValidationError as exc:
            error_rows.append({"row": row_number, "message": f"Błąd generowania kodu kreskowego: {exc}"})
            continue

        try:
            location, new_ids = resolve_location_path(root_location, row["location_path"])
            created_location_ids.update(new_ids)
        except (ValidationError, Exception) as exc:
            error_rows.append({"row": row_number, "message": f"Błąd lokalizacji: {exc}"})
            continue

        try:
            asset = Asset(
                name=row["name"],
                inventory_number=row.get("inventory_number", ""),
                asset_type_ref=asset_type_ref,
                current_quantity=quantity,
                record_quantity=quantity,
                purchase_value=purchase_value,
                location_fk=location,
                barcode=barcode,
                status=Asset.Status.ACTIVE,
            )
            asset.save()
            record_asset_history(
                asset=asset,
                operator=operator,
                event_type=AssetHistoryEntry.EventType.CREATED,
                description="Środek zaimportowany z pliku XLSX",
            )
            imported_count += 1
        except (ValidationError, Exception) as exc:
            error_rows.append({"row": row_number, "message": f"Błąd zapisu: {exc}"})

    return {
        "imported_count": imported_count,
        "created_locations_count": len(created_location_ids),
        "error_rows": error_rows,
    }


def build_import_template_xlsx():
    """Build XLSX import template. Returns BytesIO ready for HTTP response."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Szablon"

    ws.append(["Numer inwentarzowy", "Nazwa", "Ilość", "Wartość", "Lokalizacja"])
    ws.append(["INW-001", "Laptop Dell Latitude", "1", "1500.00", "Budynek A/Pokój 101"])
    ws.append(["", "Krzesło biurowe", "4", "", "Budynek A/Sala konferencyjna"])
    ws.append([])
    ws.append(["# Instrukcja:"])
    ws.append(["# Kolumna Lokalizacja: ścieżka względem wybranego roota, np. Budynek A/Pokój 1"])
    ws.append(["# Numer inwentarzowy i Wartość są opcjonalne. Ilość domyślnie = 1."])
    ws.append(["# Wiersze zaczynające się od # są ignorowane podczas importu."])

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 36

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
